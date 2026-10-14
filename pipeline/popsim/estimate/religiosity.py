"""Religiosity: P(religiosity | group, sex, age band, locality profile).

Jews (5-point scale: haredi, dati, masorti_dati, masorti, hiloni)
 1. National target by sex x age band: seed = CBS Social Survey 2024 religiosity by sex and
    5-year age group (table generator); raked (IPF) to the published Social Survey margins (sex x 4 categories, age 20-44/45-64/65+
    x 4 categories, and the 12.0 : 21.5 split of "traditional" into religious-traditional and
    not-so-religious traditional) and to the frame's Jewish adults by sex and age band.
 2. Locality profiles: each locality's religiosity mix is inferred from its 2022 vote
    (religiosity_ei.py); localities are binned by Haredi share, Dati+Masorti-dati share and
    secular share into profiles, and the profile mixes are raked to the national totals.
 3. Per profile, IPF of the national conditional pattern to the profile's adults by sex x age
    band and to the profile's religiosity totals.
Arabs and Druze (4-point scale): INES shape by religion (Muslim vs. Christian/Druze) raked to
CBS 2-level religiosity of non-Jews by sex and by age group.
Others (non-Arab, not Jewish by religion): INES "other" respondents (Jewish scale, n=30)
shrunk toward CBS figures for immigrants of 1990+ with Haredim set to zero.
"""

from __future__ import annotations

import numpy as np
import openpyxl
import pandas as pd

from .. import paths
from ..sources import load_sources
from .common import AGE_BAND_IDS, AGE_BANDS, SEXES, ipf

JEW_R = ["haredi", "dati", "masorti_dati", "masorti", "hiloni"]
ARAB_R = ["very_religious", "religious", "not_so_religious", "not_religious"]
R4_OF = {"haredi": 0, "dati": 1, "masorti_dati": 2, "masorti": 2, "hiloni": 3}  # CBS 4-category grouping
AGE3 = [(18, 44), (45, 64), (65, 200)]  # CBS age groups (20+; 18-19 joins the first)


def _age3_of_band(lo: int) -> int:
    return next(i for i, (a, b) in enumerate(AGE3) if a <= lo <= b)


def read_cbs() -> dict:
    """CBS Social Survey 2024 religiosity margins (percent)."""
    src = load_sources()["cbs_abstract_2025_attitudes"]
    f04 = src.raw_path(next(f for f in src.files if f.path == "st28_04x.xlsx"))
    f01 = src.raw_path(next(f for f in src.files if f.path == "st28_01x.xlsx"))
    ws = openpyxl.load_workbook(f04, read_only=True, data_only=True).active
    rows = {}
    for row in ws.iter_rows(values_only=True):
        if isinstance(row[0], str):
            rows[row[0].strip()] = [c for c in row[1:] if isinstance(c, (int, float)) or (isinstance(c, str) and c.replace(".", "").isdigit())]
    def jews(label):  # columns: other(2 pct, total), jews(secular, traditional, religious, haredi, total), grand total
        v = [float(x) for x in rows[label]]
        return {"other_not": v[0], "other_rel": v[1], "secular": v[3], "traditional": v[4], "dati": v[5], "haredi": v[6]}
    out = {"Men": jews("Men"), "Women": jews("Women"), "20-44": jews("20-44"), "45-64": jews("45-64"), "65+": jews("65+")}
    ws1 = openpyxl.load_workbook(f01, read_only=True, data_only=True).active
    split = {}
    for row in ws1.iter_rows(values_only=True):
        if isinstance(row[0], str) and row[0].strip() in ("Traditional-religious", "Traditional not-so-religious", "Ultra-Orthodox", "Religious", "Not religious, secular"):
            split[row[0].strip()] = float(row[1])
    out["jews_5"] = {
        "haredi": split["Ultra-Orthodox"],
        "dati": split["Religious"],
        "masorti_dati": split["Traditional-religious"],
        "masorti": split["Traditional not-so-religious"],
        "hiloni": split["Not religious, secular"],
    }
    return out


def survey_religiosity_by_age() -> pd.DataFrame:
    """P(religiosity | sex, survey age group) for Jews aged 20+, from the CBS Social Survey 2024
    table generator (religiosity x sex x 5-year age group, n ~ 5,500), shrunk toward the
    main-effects fit of sex and age (survey.main_effects)."""
    from .survey import load, main_effects

    ss = load("marital_jews")  # religiosity x sex x age group x marital status
    return main_effects(ss, "religiosity", ["sex", "age_s"], [], JEW_R, strength=10.0)


def national_jewish(frame: pd.DataFrame, cbs: dict) -> np.ndarray:
    """Counts [sex, band, r5] for Jewish adults consistent with CBS margins."""
    from .survey import AGE_BAND_TO_SURVEY

    pr = survey_religiosity_by_age().set_index(["sex", "age_s", "religiosity"])["p"]
    n = frame[frame["group"] == "jewish"].groupby(["sex", "age_band"])["count"].sum()
    seed = np.zeros((2, len(AGE_BANDS), 3, 4, 2))  # sex, band, age3, r4, within-r4
    within = {r: (1 if r == "masorti" else 0) for r in JEW_R}
    for s_i, s in enumerate(SEXES):
        for b_i, (lo, hi) in enumerate(AGE_BANDS):
            band = AGE_BAND_IDS[b_i]
            a3 = _age3_of_band(lo)
            for r in JEW_R:
                seed[s_i, b_i, a3, R4_OF[r], within[r]] = n[(s, band)] * pr[(s, AGE_BAND_TO_SURVEY[band], r)]
    total = seed.sum()
    r4_keys = ["haredi", "dati", "traditional", "secular"]
    sex_r4 = np.array([[cbs[lbl][k] for k in r4_keys] for lbl in ("Men", "Women")])
    sex_tot = seed.sum(axis=(1, 2, 3, 4))
    sex_r4 = sex_r4 / sex_r4.sum(axis=1, keepdims=True) * sex_tot[:, None]
    age_r4 = np.array([[cbs[lbl][k] for k in r4_keys] for lbl in ("20-44", "45-64", "65+")])
    age_tot = seed.sum(axis=(0, 1, 3, 4))
    age_r4 = age_r4 / age_r4.sum(axis=1, keepdims=True) * age_tot[:, None]
    j5 = cbs["jews_5"]
    trad = np.array([j5["masorti_dati"], j5["masorti"]])
    r4_within = np.zeros((4, 2))
    r4_within[:, 0] = [j5["haredi"], j5["dati"], j5["masorti_dati"], j5["hiloni"]]
    r4_within[2, 1] = j5["masorti"]
    r4_within = r4_within / r4_within.sum() * total
    del trad
    pop = seed.sum(axis=(3, 4))
    fitted = ipf(seed, [((0, 1, 2), pop), ((0, 3), sex_r4), ((2, 3), age_r4), ((3, 4), r4_within)], iters=500)
    out = np.zeros((2, len(AGE_BANDS), 5))
    for k, r in enumerate(JEW_R):
        out[:, :, k] = fitted[:, :, :, R4_OF[r], within[r]].sum(axis=2)
    return out


def locality_profiles(frame: pd.DataFrame, ei: pd.DataFrame) -> tuple[pd.Series, pd.DataFrame]:
    """locality -> profile id; profile x r5 composition (adult Jews)."""
    jew_adults = frame[frame["group"] == "jewish"].groupby("locality")["count"].sum()
    ei = ei.set_index(ei["locality"].astype(str))
    def bin_of(x: pd.Series) -> str:
        h = np.digitize(x["haredi"], [0.05, 0.15, 0.35, 0.6, 0.8])
        d = np.digitize(x["dati"] + x["masorti_dati"], [0.2, 0.3, 0.45, 0.6])
        s = np.digitize(x["hiloni"], [0.3, 0.45, 0.6])
        return f"h{h}d{d}s{s}"
    profile = {}
    for loc in jew_adults.index:
        profile[loc] = bin_of(ei.loc[loc]) if loc in ei.index else "national"
    profile = pd.Series(profile)
    comp = []
    for loc, adults in jew_adults.items():
        if loc in ei.index:
            comp.append([profile[loc]] + list(adults * ei.loc[loc, JEW_R].values.astype(float)))
    comp = pd.DataFrame(comp, columns=["profile"] + JEW_R).groupby("profile").sum()
    return profile, comp


def arab_table(ines: pd.DataFrame, cbs: dict, frame: pd.DataFrame) -> list[tuple]:
    a = ines[(ines["sector"] == "arab") & ines["religiosity"].isin(ARAB_R) & ines["w_panel1"].notna()]
    rows = []
    shapes = {}
    for cls, groups in (("muslim", ["muslim"]), ("christian_or_other", ["christian_or_other", "unknown"])):
        d = a[a["group"].isin(groups)].groupby("religiosity")["w_panel1"].sum().reindex(ARAB_R).fillna(0).values
        shapes[cls] = (d + 1) / (d.sum() + 4)
    for group in ("muslim", "christian_arab", "druze"):
        shape = shapes["muslim" if group == "muslim" else "christian_or_other"]
        n = frame[frame["group"] == group].groupby(["sex", "age_band"])["count"].sum()
        seed = np.zeros((2, len(AGE_BANDS), 3, 2, 2))  # sex, band, age3, rel2(religious/not), within
        for s_i, s in enumerate(SEXES):
            for b_i, (lo, _hi) in enumerate(AGE_BANDS):
                a3 = _age3_of_band(lo)
                c = n.get((s, AGE_BAND_IDS[b_i]), 0.0)
                seed[s_i, b_i, a3, 0, 0] = c * shape[0]
                seed[s_i, b_i, a3, 0, 1] = c * shape[1]
                seed[s_i, b_i, a3, 1, 0] = c * shape[2]
                seed[s_i, b_i, a3, 1, 1] = c * shape[3]
        tot_sex = seed.sum(axis=(1, 2, 3, 4))
        tot_age = seed.sum(axis=(0, 1, 3, 4))
        # CBS non-Jews: share religious (very religious + religious) by sex and by age group,
        # applied as relative adjustments around this group's own level
        base = shape[0] + shape[1]
        nat = (cbs["Men"]["other_rel"] + cbs["Women"]["other_rel"]) / 2
        sex_rel = np.array([cbs["Men"]["other_rel"], cbs["Women"]["other_rel"]]) / nat * base
        age_rel = np.array([cbs[k]["other_rel"] for k in ("20-44", "45-64", "65+")]) / nat * base
        sex_rel = np.clip(sex_rel, 0.02, 0.98)
        age_rel = np.clip(age_rel, 0.02, 0.98)
        m_sex = np.stack([sex_rel, 1 - sex_rel], axis=1) * tot_sex[:, None]
        m_age = np.stack([age_rel, 1 - age_rel], axis=1) * tot_age[:, None]
        fitted = ipf(seed, [((0, 1, 2), seed.sum(axis=(3, 4))), ((0, 3), m_sex)], iters=300)
        fitted = ipf(fitted, [((0, 1, 2), seed.sum(axis=(3, 4))), ((2, 3), m_age), ((3, 4), fitted.sum(axis=(0, 1, 2)))], iters=300)
        for s_i, s in enumerate(SEXES):
            for b_i, band in enumerate(AGE_BAND_IDS):
                cell = fitted[s_i, b_i].sum(axis=0).reshape(-1)  # [rel x within] -> 4 values in ARAB_R order
                if cell.sum() <= 0:
                    cell = shape.copy()
                cell = cell / cell.sum()
                for k, r in enumerate(ARAB_R):
                    rows.append((group, s, band, "*", r, cell[k]))
    return rows


def other_table(ines: pd.DataFrame) -> list[tuple]:
    o = ines[(ines["group"] == "other") & ines["religiosity"].isin(JEW_R)]
    d = o.groupby("religiosity")["w_panel1"].sum().reindex(JEW_R).fillna(0)
    # CBS immigrants of 1990+ (Social Survey 2024): secular 54.5, traditional 29.1, religious 10,
    # Haredi 6.2 -> Haredim set to zero; traditional split like all Jews (12.0 : 21.5)
    prior = pd.Series({"haredi": 0.0, "dati": 10.0, "masorti_dati": 29.1 * 12 / 33.5, "masorti": 29.1 * 21.5 / 33.5, "hiloni": 54.5})
    prior = prior / prior.sum()
    n = d.sum()
    p = (d / max(n, 1e-9) * n + prior * 30) / (n + 30)
    rows = []
    for s in SEXES:
        for band in AGE_BAND_IDS:
            for r in JEW_R:
                rows.append(("other", s, band, "*", r, float(p[r])))
    return rows


def build(log=print) -> None:
    ines = pd.read_csv(paths.interim_path("ines_2022"))
    frame = pd.read_csv(paths.MODEL / "frame.csv", dtype={"locality": str})
    ei = pd.read_csv(paths.interim_path("religiosity_by_locality"))
    cbs = read_cbs()

    national = national_jewish(frame, cbs)  # [sex, band, r]
    profile_of, comp = locality_profiles(frame, ei)
    # rake profile compositions so they add up to the national totals
    prof_adults = frame[frame["group"] == "jewish"].assign(profile=lambda d: d["locality"].map(profile_of)).groupby("profile")["count"].sum()
    comp = comp.reindex(prof_adults.index).fillna(0)
    if "national" in comp.index:  # localities without a vote signal follow the national mix
        comp.loc["national"] = national.sum(axis=(0, 1)) / national.sum() * prof_adults["national"]
    fitted_comp = ipf(comp.values, [((0,), prof_adults.values), ((1,), national.sum(axis=(0, 1)) / national.sum() * prof_adults.sum())])
    comp = pd.DataFrame(fitted_comp, index=comp.index, columns=JEW_R)

    nat_cond = national / national.sum(axis=2, keepdims=True)
    rows = []
    jf = frame[frame["group"] == "jewish"].assign(profile=lambda d: d["locality"].map(profile_of))
    for prof, g in jf.groupby("profile"):
        n = g.groupby(["sex", "age_band"])["count"].sum()
        sb = np.array([[n.get((s, b), 0.0) for b in AGE_BAND_IDS] for s in SEXES])
        seed = nat_cond * sb[:, :, None]
        fitted = ipf(seed, [((0, 1), sb), ((2,), comp.loc[prof].values * sb.sum() / max(comp.loc[prof].sum(), 1e-9))])
        for s_i, s in enumerate(SEXES):
            for b_i, band in enumerate(AGE_BAND_IDS):
                cell = fitted[s_i, b_i]
                cell = cell / cell.sum() if cell.sum() > 0 else nat_cond[s_i, b_i]
                for k, r in enumerate(JEW_R):
                    rows.append(("jewish", s, band, prof, r, float(cell[k])))
    # Jews in a profile/sex/band combination with no adults in the frame never occur; a national row backs them off
    for s_i, s in enumerate(SEXES):
        for b_i, band in enumerate(AGE_BAND_IDS):
            for k, r in enumerate(JEW_R):
                rows.append(("jewish", s, band, "*", r, float(nat_cond[s_i, b_i, k])))
    rows += arab_table(ines, cbs, frame)
    rows += other_table(ines)
    cpt = pd.DataFrame(rows, columns=["group", "sex", "age_band", "locality_profile", "value", "p"])
    cpt.to_csv(paths.CPTS / "religiosity.csv", index=False)

    # locality_profile lookup (hidden node): every locality gets a profile
    loc = pd.read_csv(paths.MODEL / "domains" / "locality.csv", dtype={"id": str})
    lk = pd.DataFrame({"locality": loc["id"], "value": loc["id"].map(profile_of).fillna("national")})
    lk.to_csv(paths.MODEL / "lookups" / "locality_profile.csv", index=False)
    profiles = sorted(set(lk["value"]))
    pd.DataFrame({"id": profiles, "en": profiles, "he": profiles}).to_csv(paths.MODEL / "domains" / "locality_profile.csv", index=False)
    tot = national.sum(axis=(0, 1))
    log(f"  religiosity: {len(profiles)} profiles; national Jewish mix " + ", ".join(f"{r} {v / tot.sum():.3f}" for r, v in zip(JEW_R, tot)))
