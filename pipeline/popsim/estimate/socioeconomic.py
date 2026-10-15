"""Family, education, service, work and income nodes from CBS Social Survey 2024 tables.

Nodes: birth_region (lookup), marital_status, children, has_children (lookup),
living_arrangement, education, military_service,
in_service, labour_status, occupation, industry, work_status, income, income_amount,
income_percentile, household_income, home_ownership. Each has a methodology page in
docs/methodology/.

Common pattern. Survey tables (persons aged 20+, data/manual/cbs_ss/) come per population:
Jews by religiosity, "others" (neither Jewish nor Arab) and Arabs by religion - the Arab
tables exclude the Jerusalem district, whose Arab residents are mostly not citizens. Each
conditional is shrunk toward a main-effects fit (survey.main_effects) and lightly smoothed
across the 45-64 age groups. When two published tables condition on different
variables (employment by religiosity, and by education), they are joined with
common.combine(): iterative proportional fitting over the model's own population counts
(joint.joint). 18-19-year-olds, whom the survey does not cover, use the 20-24 pattern unless
a node says otherwise.
"""

from __future__ import annotations

from collections.abc import Callable

import numpy as np
import openpyxl
import pandas as pd

from .. import paths
from ..model_spec import load_spec
from ..sources import load_sources
from .common import AGE_BAND_IDS, combine, normalize_rows
from .joint import joint
from .survey import AGE_BAND_TO_SURVEY, load, main_effects

FAMILY = {"jewish": "jews", "other": "others", "muslim": "arabs", "christian_arab": "arabs", "druze": "arabs"}
ARAB_RELIGION = {"muslim": "muslim", "christian_arab": "christian", "druze": "druze"}
POPGROUP = {"jewish": "jews_and_others", "other": "jews_and_others", "muslim": "arabs", "christian_arab": "arabs", "druze": "arabs"}
SURVEY_AGES = ["20-24", "25-29", "30-34", "35-39", "40-44", "45-49", "50-54", "55-59", "60-64", "65-74", "75+"]
SMOOTHED_AGES = ["45-49", "50-54", "55-59", "60-64"]
BANDS = pd.DataFrame({"age_band": AGE_BAND_IDS, "age_s": [AGE_BAND_TO_SURVEY[a] for a in AGE_BAND_IDS]})
OLD_BANDS = ["65-69", "70-74", "75-79", "80-84", "85+"]
STRENGTH = 10.0  # pseudo-respondents of the main-effects prior (survey.main_effects)


# ---------------------------------------------------------------------------
# shared helpers


def smooth_ages(df: pd.DataFrame, target: str) -> pd.DataFrame:
    """Damp sampling noise in the 45-64 age groups: each becomes 1/4 previous + 1/2 itself +
    1/4 next group (a kernel that keeps linear age trends). Other groups are unchanged."""
    keys = [c for c in df.columns if c not in (target, "p", "n_eff", "age_s")]
    piv = df.pivot_table(index=keys + [target], columns="age_s", values="p", aggfunc="sum").reindex(columns=SURVEY_AGES)
    out = piv.copy()
    for a in SMOOTHED_AGES:
        i = SURVEY_AGES.index(a)
        out[a] = 0.25 * piv[SURVEY_AGES[i - 1]] + 0.5 * piv[a] + 0.25 * piv[SURVEY_AGES[i + 1]]
    long = out.stack(future_stack=True).rename("p").reset_index()
    long = long[long["p"].notna()]
    long["p"] = long["p"] / long.groupby(keys + ["age_s"])["p"].transform("sum")
    return long


def by_population(name: str, target: str, values: list[str], extra: list[str] | None = None, loader: Callable[[str], pd.DataFrame] = load) -> dict[str, pd.DataFrame]:
    """P(target | ...) for Jews by religiosity, others, and Arabs by religion (x sex x age group).

    Others join the Jewish table as one more "religiosity" level, so their few respondents
    borrow the age pattern of all Jews-and-others through the main-effects prior.
    """
    by = ["sex"] + (extra or [])
    j = loader(f"{name}_jews")
    o = loader(f"{name}_others").assign(religiosity="other")
    a = loader(f"{name}_arabs")
    a = a[a["religion"].isin(["muslim", "christian", "druze"])]
    jo = main_effects(pd.concat([j, o], ignore_index=True), target, ["religiosity", "age_s"], by, values, STRENGTH)
    out = {
        "jews": jo[jo["religiosity"] != "other"],
        "others": jo[jo["religiosity"] == "other"].drop(columns=["religiosity"]),
        "arabs": main_effects(a, target, ["religion", "age_s"], by, values, STRENGTH),
    }
    return {k: smooth_ages(v, target) for k, v in out.items()}


def expand(cond: dict[str, pd.DataFrame], target: str, extra: list[str] | None = None) -> pd.DataFrame:
    """Survey conditionals -> model rows [group, religiosity, sex, age_band, *extra, value, p].
    Non-Jewish groups get religiosity "*" (their rows do not depend on it)."""
    extra = extra or []
    parts = [cond["jews"].merge(BANDS, on="age_s").assign(group="jewish")]
    parts.append(cond["others"].merge(BANDS, on="age_s").assign(group="other", religiosity="*"))
    arabs = cond["arabs"].merge(BANDS, on="age_s")
    for g, rel in ARAB_RELIGION.items():
        parts.append(arabs[arabs["religion"] == rel].assign(group=g, religiosity="*"))
    df = pd.concat(parts, ignore_index=True).rename(columns={target: "value"})
    return df[["group", "religiosity", "sex", "age_band"] + extra + ["value", "p"]]


def write_cpt(df: pd.DataFrame, name: str) -> pd.DataFrame:
    df = df[df["p"] > 1e-12].copy()
    df["p"] = df["p"].round(10)
    df.to_csv(paths.CPTS / f"{name}.csv", index=False)
    return df


def _rows_for(df: pd.DataFrame, **eq) -> pd.DataFrame:
    m = np.ones(len(df), dtype=bool)
    for k, v in eq.items():
        m &= (df[k] == v).to_numpy()
    return df[m]


def combine_by_cell(W: pd.DataFrame, a: str, b: str, A: pd.DataFrame, B: pd.DataFrame, target: str, values: list[str]) -> pd.DataFrame:
    """common.combine() separately for every (sex, age band). A and B are keyed by survey age
    group (age_s); W by model age band."""
    out = []
    for (sex, band), w in W.groupby(["sex", "age_band"]):
        age_s = AGE_BAND_TO_SURVEY[band]
        Ac = _rows_for(A, sex=sex, age_s=age_s)
        Bc = _rows_for(B, sex=sex, age_s=age_s)
        c = combine(w[[a, b, "w"]], a, b, Ac[[a, target, "p"]], Bc[[b, target, "p"]], target, values)
        out.append(c.assign(sex=sex, age_band=band))
    return pd.concat(out, ignore_index=True)


# ---------------------------------------------------------------------------
# birth region (lookup from origin)

AFRICA = {"morocco", "ethiopia", "algeria_tunisia", "libya", "egypt", "africa_other"}
ASIA = {"iraq", "iran", "yemen", "turkey", "india", "syria_lebanon", "asia_other"}


def build_birth_region(log=print) -> None:
    """Continent of birth, as the CBS groups it (the former USSR counts as Europe)."""
    origins = load_spec().node("origin").value_ids()
    rows = []
    for o in origins:
        if o.endswith("_immigrant"):
            c = o[: -len("_immigrant")]
            rows.append((o, "africa" if c in AFRICA else "asia" if c in ASIA else "europe_america"))
        else:
            rows.append((o, "israel"))
    pd.DataFrame(rows, columns=["origin", "value"]).to_csv(paths.MODEL / "lookups" / "birth_region.csv", index=False)
    log("  birth_region")


# ---------------------------------------------------------------------------
# marital status and children

MARITAL = ["married", "divorced", "widowed", "never_married"]
ABSTRACT_RELIGION = {"jewish": "jews", "muslim": "muslims", "christian_arab": "christians", "druze": "druze"}


def married_18_19() -> pd.DataFrame:
    """Share married at ages 18-19 by religion and sex, from marriage flows.

    In a steady state, the married stock at 18-19 is about 2 m(<=17) + 1.5 m(18) + 0.5 m(19)
    (annual marriages by age at marriage; someone marrying at 18.5 spends 1.5 more years in
    the band). Divided by the population aged 18-19 (Statistical Abstract table 2.3).
    """
    src = load_sources()["cbs_abstract_2025_marriages"]
    ws = openpyxl.load_workbook(src.raw_path(src.files[0]), read_only=True, data_only=True).worksheets[0]
    # total columns (all previous statuses): Druze 1, Christians 3, Muslims 5, Jews 8
    cols = {"druze": 1, "christians": 3, "muslims": 5, "jews": 8}
    sex = None
    flows: dict[tuple[str, str], dict[str, float]] = {}
    for row in ws.iter_rows(values_only=True):
        first = str(row[0] or "")
        label = str(row[11]).strip() if len(row) > 11 and row[11] is not None else ""  # age column
        if "GROOMS" in first:
            sex = "male"
            continue
        if "BRIDES" in first:
            sex = "female"
            continue
        if sex is None:
            continue
        age = {"18": "18", "19": "19"}.get(label.removesuffix(".0"))
        if "עד 19" in label:
            age = "le19"
        if age is None:
            continue
        for rel, c in cols.items():
            v = row[c]
            flows.setdefault((rel, sex), {})[age] = float(v) if isinstance(v, (int, float)) else 0.0
    pop = pd.read_csv(paths.interim_path("cbs_pop_age"))
    pop = pop[(pop["age_from"] == pop["age_to"]) & pop["age_from"].isin([18, 19])]
    rows = []
    for (rel, sx), f in flows.items():
        m17 = f["le19"] - f["18"] - f["19"]
        stock = 2 * m17 + 1.5 * f["18"] + 0.5 * f["19"]
        n = pop[(pop["population"] == rel) & (pop["sex"] == sx)]["thousands"].sum() * 1000
        rows.append((rel, sx, stock / n))
    return pd.DataFrame(rows, columns=["religion", "sex", "p_married"])


def _scale_odds(p: np.ndarray, w: np.ndarray, target: float) -> np.ndarray:
    """Multiply the odds of each p by one factor so that the w-weighted mean equals target."""
    p = np.clip(p, 1e-9, 1 - 1e-9)
    lo, hi = -30.0, 30.0
    for _ in range(200):
        mid = (lo + hi) / 2
        q = 1 / (1 + (1 - p) / p * np.exp(-mid))
        if np.average(q, weights=w) > target:
            hi = mid
        else:
            lo = mid
    return 1 / (1 + (1 - p) / p * np.exp(-(lo + hi) / 2))


def build_marital(log=print) -> None:
    df = expand(by_population("marital", "marital", MARITAL), "marital")
    # 18-19: scale the odds of being married (20-24 pattern) to the level implied by marriage flows
    target = married_18_19()
    rel_w = joint(["group", "sex", "religiosity"], where={"age_band": ["18-19"]})
    young = df[df["age_band"] == "18-19"]
    rest = df[df["age_band"] != "18-19"]
    out = []
    for (group, sex), g in young.groupby(["group", "sex"]):
        rel = ABSTRACT_RELIGION.get(group, "jews")  # "others" follow the Jewish level
        t = float(target[(target["religion"] == rel) & (target["sex"] == sex)]["p_married"].iloc[0])
        piv = g.pivot_table(index="religiosity", columns="value", values="p").reindex(columns=MARITAL).fillna(0.0)
        if group == "jewish":
            w = rel_w[(rel_w["group"] == group) & (rel_w["sex"] == sex)].set_index("religiosity")["w"].reindex(piv.index).fillna(0.0).to_numpy()
        else:
            w = np.ones(len(piv))
        pm = _scale_odds(piv["married"].to_numpy(), w, t)
        ratio = pm / np.clip(piv["married"].to_numpy(), 1e-9, None)
        new = piv.copy()
        new["married"] = pm
        new["divorced"] = np.minimum(piv["divorced"].to_numpy() * ratio, piv["divorced"].to_numpy())
        new["widowed"] = 0.0
        new["never_married"] = 1 - new[["married", "divorced", "widowed"]].sum(axis=1)
        new.columns.name = "value"
        long = new.stack().rename("p").reset_index()
        out.append(long.assign(group=group, sex=sex, age_band="18-19"))
    df = pd.concat([rest] + out, ignore_index=True)[["group", "religiosity", "sex", "age_band", "value", "p"]]
    write_cpt(df, "marital_status")
    log(f"  marital_status (married at 18-19: {', '.join(f'{r.religion}/{r.sex} {r.p_married:.3f}' for r in target.itertuples())})")


CHILDREN = [str(i) for i in range(8)]


def build_children(log=print) -> None:
    """Children ever born. 18-19-year-olds use the 20-24 pattern for their marital status,
    capped at one child."""
    cond = by_population("children", "children", CHILDREN, extra=["marital"])
    df = expand(cond, "children", extra=["marital"]).rename(columns={"marital": "marital_status"})
    keys = ["group", "religiosity", "sex", "marital_status"]
    young = df["age_band"] == "18-19"
    many = young & df["value"].isin(CHILDREN[2:])
    moved = df[many].groupby(keys)["p"].sum().rename("moved").reset_index()
    df.loc[many, "p"] = 0.0
    one = df[young & (df["value"] == "1")].reset_index().merge(moved, on=keys, how="left").set_index("index")
    df.loc[one.index, "p"] += one["moved"].fillna(0.0)
    write_cpt(df[["group", "religiosity", "sex", "age_band", "marital_status", "value", "p"]], "children")
    pd.DataFrame({"children": CHILDREN, "value": ["no"] + ["yes"] * (len(CHILDREN) - 1)}).to_csv(paths.MODEL / "lookups" / "has_children.csv", index=False)
    log("  children, has_children")


# ---------------------------------------------------------------------------
# living arrangement

LIVING = ["partner", "parents", "own_children", "relatives", "independent"]
OLD_SURVEY_AGES = ["65-74", "75+"]


def arrangement(df: pd.DataFrame) -> pd.Series:
    """Living arrangement of a Social Survey respondent, from the household type, whether the
    partner lives in the household, whether the respondent has had children, and age.

    A CBS household shares a dwelling and a food budget, so flatmates who buy food separately
    are one-person households: "independent" is a one-person or non-family household. A CBS
    family is a couple, a parent and child, or siblings without partners or children: a
    respondent without a partner in a family household lives with their own children if they
    have had any, otherwise with their family of origin - parents before 65, other relatives
    (mostly siblings) from 65 on.
    """
    old = df["age_s"].isin(OLD_SURVEY_AGES)
    return pd.Series(
        np.select(
            [df["with_partner"] == "yes", df["hh_type"].isin(["one_person", "non_family"]), df["had_children"] == "yes", old],
            ["partner", "independent", "own_children", "relatives"],
            "parents",
        ),
        index=df.index,
    )


def load_living(name: str) -> pd.DataFrame:
    """<name>_children and <name>_childless (tables filtered on whether the respondent has had
    children), with the respondent's living arrangement."""
    parts = [load(f"{name}_{k}").assign(had_children=v) for k, v in (("children", "yes"), ("childless", "no"))]
    df = pd.concat(parts, ignore_index=True)
    return df.assign(arrangement=arrangement(df)).rename(columns={"had_children": "has_children"})


def impossible_living(df: pd.DataFrame, age: str) -> np.ndarray:
    """Arrangements ruled out by definition: own children without having had any; parents
    from 65 on; other relatives before 65."""
    old = df[age].isin(OLD_SURVEY_AGES if age == "age_s" else OLD_BANDS)
    v = df["value"] if "value" in df else df["arrangement"]
    return (((v == "own_children") & (df["has_children"] == "no")) | ((v == "parents") & old) | ((v == "relatives") & ~old)).to_numpy()


def build_living(log=print) -> None:
    """Living arrangement given marital status and whether the person has had children, by
    religiosity (Jews) or religion (Arabs), sex and age. 18-19-year-olds use the 20-24 pattern.

    Each (marital status, children) slice is fitted on its own, so that a row without
    respondents (never-married Haredi women aged 30-34) borrows from people in the same
    family situation - among whom few live with a partner - not from the married."""
    tables = {pop: load_living(f"living_{pop}") for pop in ("jews", "others", "arabs")}
    parts = []
    for (marital, kids), _ in tables["jews"].groupby(["marital", "has_children"]):
        sl = {pop: t[(t["marital"] == marital) & (t["has_children"] == kids)] for pop, t in tables.items()}
        cond = by_population("living", "arrangement", LIVING, loader=lambda name: sl[name.removeprefix("living_")])
        parts.append(expand(cond, "arrangement").assign(marital_status=marital, has_children=kids))
    df = pd.concat(parts, ignore_index=True)
    df.loc[impossible_living(df, "age_band"), "p"] = 0.0
    df = normalize_rows(df, ["group", "religiosity", "sex", "age_band", "marital_status", "has_children"])
    write_cpt(df[["group", "religiosity", "sex", "age_band", "marital_status", "has_children", "value", "p"]], "living_arrangement")
    log("  living_arrangement")


# ---------------------------------------------------------------------------
# education

EDUCATION = ["no_diploma", "primary", "secondary", "bagrut", "post_secondary", "ba", "ma", "phd", "other_diploma"]


def _birth_conditional(name: str, target: str, values: list[str], loader: Callable[[str], pd.DataFrame] = load) -> pd.DataFrame:
    """P(target | continent of birth, sex, age group) for Jews (the B tables of combine())."""
    b = loader(name).rename(columns={"birth_continent": "birth_region"})
    return smooth_ages(main_effects(b, target, ["birth_region", "age_s"], ["sex"], values, STRENGTH), target)


def _with_birth(A: dict[str, pd.DataFrame], B: pd.DataFrame, target: str, values: list[str]) -> pd.DataFrame:
    """Jews and others: combine P(t | religiosity) with P(t | birth region) over the model's
    population counts. Returns [group, religiosity, sex, age_band, birth_region, value, p]."""
    Wj = joint(["sex", "age_band", "religiosity", "birth_region"], where={"group": ["jewish"]})
    cj = combine_by_cell(Wj, "religiosity", "birth_region", A["jews"], B, target, values).assign(group="jewish")
    Wo = joint(["sex", "age_band", "birth_region"], where={"group": ["other"]}).assign(religiosity="*")
    co = combine_by_cell(Wo, "religiosity", "birth_region", A["others"].assign(religiosity="*"), B, target, values).assign(group="other")
    df = pd.concat([cj, co], ignore_index=True).rename(columns={target: "value"})
    return df[["group", "religiosity", "sex", "age_band", "birth_region", "value", "p"]]


def _arab_rows(A: dict[str, pd.DataFrame], target: str) -> pd.DataFrame:
    a = A["arabs"].merge(BANDS, on="age_s")
    parts = [a[a["religion"] == rel].assign(group=g, religiosity="*", birth_region="israel") for g, rel in ARAB_RELIGION.items()]
    return pd.concat(parts, ignore_index=True).rename(columns={target: "value"})[["group", "religiosity", "sex", "age_band", "birth_region", "value", "p"]]


def build_education(log=print) -> None:
    """Highest diploma, by religiosity (or religion), sex, age and continent of birth.
    18-19-year-olds: the 20-24 pattern with degrees and post-secondary diplomas moved to
    matriculation (bagrut), since they are too young to hold them."""
    A = by_population("diploma", "diploma", EDUCATION)
    B = _birth_conditional("diploma_by_birth_jews", "diploma", EDUCATION)
    df = pd.concat([_with_birth(A, B, "diploma", EDUCATION), _arab_rows(A, "diploma")], ignore_index=True)
    keys = ["group", "religiosity", "sex", "birth_region"]
    young = df["age_band"] == "18-19"
    adv = young & df["value"].isin(["post_secondary", "ba", "ma", "phd"])
    moved = df[adv].groupby(keys)["p"].sum().rename("moved").reset_index()
    df.loc[adv, "p"] = 0.0
    bag = df[young & (df["value"] == "bagrut")].reset_index().merge(moved, on=keys, how="left").set_index("index")
    df.loc[bag.index, "p"] += bag["moved"].fillna(0.0)
    write_cpt(df, "education")
    log("  education")


# ---------------------------------------------------------------------------
# military / national service

SERVICE = ["idf", "national_service", "none"]


def service_table(name: str) -> pd.DataFrame:
    """IDF service and national service as one variable. The survey asks about national
    service only those who did not serve in the IDF, so: idf = served in the IDF;
    national_service = did not, but did national service; none = neither.
    name: "service_<suffix>" reads military_<suffix> and national_service_<suffix>."""
    suffix = name[len("service_"):]
    m = load(f"military_{suffix}")
    n = load(f"national_service_{suffix}")
    keys = [c for c in m.columns if c not in ("count", "rse", "idf")]
    idf = m[m["idf"] == "yes"].drop(columns=["idf"]).assign(service="idf")
    ns = n.assign(service=n["national_service"].map({"yes": "national_service", "no": "none"})).drop(columns=["national_service"])
    out = pd.concat([idf, ns], ignore_index=True)
    return out.groupby(keys + ["service"], as_index=False).agg(count=("count", "sum"), rse=("rse", "max"))


def build_military(log=print) -> None:
    A = by_population("service", "service", SERVICE, loader=service_table)
    B = _birth_conditional("service_by_birth_jews", "service", SERVICE, loader=service_table)
    df = pd.concat([_with_birth(A, B, "service", SERVICE), _arab_rows(A, "service")], ignore_index=True)
    write_cpt(df, "military_service")

    # currently serving: men aged 18-20 and women aged 18-19 in the IDF; national service at 18-19
    rows = [("*", "*", "*", "no", 1.0)]
    for age in (18, 19, 20):
        for sex, last in (("male", 20), ("female", 19)):
            if age <= last:
                rows.append((str(age), sex, "idf", "idf", 1.0))
            if age <= 19:
                rows.append((str(age), sex, "national_service", "national_service", 1.0))
    pd.DataFrame(rows, columns=["age", "sex", "military_service", "value", "p"]).to_csv(paths.CPTS / "in_service.csv", index=False)
    log("  military_service, in_service")


# ---------------------------------------------------------------------------
# labour force status, occupation, industry, status at work

EMPLOYMENT = ["employed", "unemployed", "not_in_labour_force"]


def with_second(A: dict[str, pd.DataFrame], B: dict[str, pd.DataFrame] | pd.DataFrame, b: str, target: str, values: list[str],
                W: Callable[[list[str], dict], pd.DataFrame] | None = None) -> pd.DataFrame:
    """Combine P(t | religiosity or religion, sex, age) with P(t | b, sex, age) - b a model
    variable such as education - over the model's joint counts of religiosity x b.
    B: one table for everyone, or {"jews_and_others": ..., "arabs": ...}. W: how to get the
    model's counts (default joint.joint; override for derived b). Returns
    [group, religiosity, sex, age_band, b, value, p]."""
    W = W or (lambda keep, where: joint(keep, where=where))
    b_jo = B["jews_and_others"] if isinstance(B, dict) else B
    b_ar = B["arabs"] if isinstance(B, dict) else B
    out = []
    Wj = W(["sex", "age_band", "religiosity", b], {"group": ["jewish"]})
    out.append(combine_by_cell(Wj, "religiosity", b, A["jews"], b_jo, target, values).assign(group="jewish"))
    Wo = W(["sex", "age_band", b], {"group": ["other"]}).assign(religiosity="*")
    out.append(combine_by_cell(Wo, "religiosity", b, A["others"].assign(religiosity="*"), b_jo, target, values).assign(group="other"))
    Wa = W(["group", "sex", "age_band", b], {"group": list(ARAB_RELIGION)})
    Aa = A["arabs"].assign(group=A["arabs"]["religion"].map({v: k for k, v in ARAB_RELIGION.items()}))
    out.append(combine_by_cell(Wa, "group", b, Aa, b_ar, target, values).assign(religiosity="*"))
    df = pd.concat(out, ignore_index=True).rename(columns={target: "value"})
    return df[["group", "religiosity", "sex", "age_band", b, "value", "p"]]


def by_diploma(name: str, target: str, values: list[str]) -> pd.DataFrame:
    d = load(name).rename(columns={"diploma": "education"})
    return smooth_ages(main_effects(d, target, ["education", "age_s"], ["sex"], values, STRENGTH), target)


def build_labour(log=print) -> None:
    """Labour-force status. People in compulsory service are "compulsory_service"; people
    aged 65+ outside the labour force are "retired"."""
    A = by_population("employment", "employment", EMPLOYMENT)
    B = {"jews_and_others": by_diploma("employment_by_diploma_jo", "employment", EMPLOYMENT),
         "arabs": by_diploma("employment_by_diploma_arabs", "employment", EMPLOYMENT)}
    df = with_second(A, B, "education", "employment", EMPLOYMENT)
    retired = df["age_band"].isin(OLD_BANDS) & (df["value"] == "not_in_labour_force")
    df.loc[retired, "value"] = "retired"
    df = df.assign(in_service="no")
    serving = pd.DataFrame([(s, "*", "*", "*", "*", "*", "compulsory_service", 1.0) for s in ("idf", "national_service")],
                           columns=["in_service", "group", "religiosity", "sex", "age_band", "education", "value", "p"])
    df = pd.concat([df, serving], ignore_index=True)[["in_service", "group", "religiosity", "sex", "age_band", "education", "value", "p"]]
    write_cpt(df, "labour_status")
    log("  labour_status")


OCCUPATION = ["managers", "professionals", "technicians", "clerical", "sales_services", "skilled_workers", "unskilled"]


def build_occupation(log=print) -> None:
    """Occupation (major group) of the employed, by religiosity (or religion), sex and
    education. Retirees get a former occupation from the same distribution; everyone else
    "none"."""
    j = load("occupation_jews").rename(columns={"diploma": "education"})
    o = load("occupation_others").rename(columns={"diploma": "education"})
    a = load("occupation_arabs").rename(columns={"diploma": "education"})
    a = a[a["religion"].isin(ARAB_RELIGION.values())]
    jo = main_effects(pd.concat([j, o.assign(religiosity="other")], ignore_index=True), "occupation", ["religiosity", "education"], ["sex"], OCCUPATION, STRENGTH)
    hj = jo[jo["religiosity"] != "other"]
    ho = jo[jo["religiosity"] == "other"].assign(religiosity="*")
    ha = main_effects(a, "occupation", ["religion", "education"], ["sex"], OCCUPATION, STRENGTH)
    parts = [hj.assign(group="jewish"), ho.assign(group="other")]
    for g, rel in ARAB_RELIGION.items():
        parts.append(ha[ha["religion"] == rel].assign(group=g, religiosity="*"))
    occ = pd.concat(parts, ignore_index=True).rename(columns={"occupation": "value"})
    occ = occ[["group", "religiosity", "sex", "education", "value", "p"]]
    rows = [occ.assign(labour_status=s) for s in ("employed", "retired")]
    rows.append(pd.DataFrame([("*", "*", "*", "*", "*", "none", 1.0)], columns=["labour_status", "group", "religiosity", "sex", "education", "value", "p"]))
    df = pd.concat(rows, ignore_index=True)[["labour_status", "group", "religiosity", "sex", "education", "value", "p"]]
    write_cpt(df, "occupation")
    log("  occupation")


INDUSTRY = ["agriculture", "manufacturing", "utilities", "construction", "trade", "transport", "hospitality", "ict", "finance",
            "real_estate", "professional", "admin_support", "public_admin", "education", "health_welfare", "arts_other",
            "households", "extraterritorial"]
WORK_STATUS = ["employee", "self_employed", "company_owner", "unpaid_family"]


def _per_popgroup(name: str) -> dict[str, pd.DataFrame]:
    """A table queried separately for Jews-and-others (<name>_jo) and Arabs outside the
    Jerusalem district (<name>_arabs)."""
    return {"jews_and_others": load(f"{name}_jo"), "arabs": load(f"{name}_arabs")}


def _per_popgroup_living(name: str) -> dict[str, pd.DataFrame]:
    """_per_popgroup() for tables that also give the household type and partner, queried
    separately for people who have and have not had children (see load_living)."""
    return {"jews_and_others": load_living(f"{name}_jo"), "arabs": load_living(f"{name}_arabs")}


def _to_groups(tables: dict[str, pd.DataFrame]) -> pd.DataFrame:
    """Copy each population-group table to the model groups it covers."""
    return pd.concat([tables[pg].assign(group=g) for g, pg in POPGROUP.items()], ignore_index=True)


def build_industry_work(log=print) -> None:
    """Industry given occupation and sex; status at work (employee, self-employed...) given
    occupation and sex. Both separately for Jews-and-others and for Arabs."""
    hi = {pg: main_effects(t, "industry", ["occupation", "sex"], [], INDUSTRY, STRENGTH)
          for pg, t in _per_popgroup("industry_by_occupation").items()}
    df = _to_groups(hi).rename(columns={"industry": "value"})[["occupation", "sex", "group", "value", "p"]]
    df = pd.concat([df, pd.DataFrame([("none", "*", "*", "none", 1.0)], columns=["occupation", "sex", "group", "value", "p"])], ignore_index=True)
    write_cpt(df, "industry")

    hw = {}
    for pg, ws in _per_popgroup("work_status_by_occupation").items():
        # kibbutz members are counted as employees (the model does not tie them to kibbutzim)
        ws["work_status"] = ws["work_status"].replace({"kibbutz_member": "employee", "cooperative_member": "employee"})
        ws = ws.groupby(["occupation", "sex", "work_status"], as_index=False).agg(count=("count", "sum"), rse=("rse", "max"))
        hw[pg] = main_effects(ws, "work_status", ["occupation", "sex"], [], WORK_STATUS, STRENGTH)
    df = _to_groups(hw).rename(columns={"work_status": "value"}).assign(labour_status="employed")
    df = df[["labour_status", "occupation", "sex", "group", "value", "p"]]
    df = pd.concat([df, pd.DataFrame([("*", "*", "*", "*", "none", 1.0)], columns=["labour_status", "occupation", "sex", "group", "value", "p"])], ignore_index=True)
    write_cpt(df, "work_status")
    log("  industry, work_status")


# ---------------------------------------------------------------------------
# earnings and household income

EARNINGS = [f"b{i:02d}" for i in range(1, 11)]
INCOME = ["not_employed"] + EARNINGS
# Gross NIS per month. hi=None: open top bracket with a Pareto tail (dag.yaml tail_alpha).
INCOME_BOUNDS: dict[str, tuple[float, float | None] | None] = {
    "not_employed": None, "b01": (500, 2000), "b02": (2001, 3000), "b03": (3001, 4000), "b04": (4001, 5000),
    "b05": (5001, 6000), "b06": (6001, 7500), "b07": (7501, 10000), "b08": (10001, 14000),
    "b09": (14001, 21000), "b10": (21001, None),
}
HH_INCOME = ["none"] + [f"h{i:02d}" for i in range(1, 11)]
HH_BOUNDS: dict[str, tuple[float, float | None] | None] = {
    "none": (0, 0), "h01": (1000, 2500), "h02": (2501, 4000), "h03": (4001, 5000), "h04": (5001, 6500),
    "h05": (6501, 8000), "h06": (8001, 10000), "h07": (10001, 13000), "h08": (13001, 17000),
    "h09": (17001, 24000), "h10": (24001, None),
}


def _write_bounds(bounds: dict, parent: str, name: str) -> None:
    rows = [(k, "", "") if b is None else (k, b[0], "" if b[1] is None else b[1]) for k, b in bounds.items()]
    pd.DataFrame(rows, columns=[parent, "lo", "hi"]).to_csv(paths.CPTS / f"{name}.csv", index=False)


def build_income(log=print) -> None:
    """Gross monthly earnings bracket of the employed, by occupation, sex and age (the survey
    asks the income question of employed persons only). Everyone else: "not_employed"."""
    rows = []
    for pg, t in _per_popgroup("income_by_occupation").items():
        h = main_effects(t, "income", ["occupation", "age_s"], ["sex"], EARNINGS, STRENGTH)
        h = smooth_ages(h, "income").merge(BANDS, on="age_s")
        for g, pg_ in POPGROUP.items():
            if pg_ == pg:
                rows.append(h.assign(group=g, labour_status="employed"))
    df = pd.concat(rows, ignore_index=True).rename(columns={"income": "value"})
    df = df[["labour_status", "group", "occupation", "sex", "age_band", "value", "p"]]
    other = [(s, "*", "*", "*", "*", "not_employed", 1.0) for s in ("unemployed", "not_in_labour_force", "retired", "compulsory_service")]
    df = pd.concat([df, pd.DataFrame(other, columns=df.columns)], ignore_index=True)
    write_cpt(df, "income")
    _write_bounds(INCOME_BOUNDS, "income", "income_amount")
    log("  income, income_amount")


def _tail(node_id: str) -> tuple[float, float]:
    raw = load_spec().node(node_id).raw
    return float(raw["tail_alpha"]), float(raw["tail_u_max"])


def build_income_percentile(log=print) -> None:
    """Earnings at percentiles 0..100 among employed citizens in the model, so a person's
    percentile is their rank among the simulated employed population. Within a bracket the
    amount is capped at the top of the person's household bracket, so each bracket's CDF is the
    mixture over household brackets of the capped distribution, inverted numerically."""
    a_in, u_in = _tail("income_amount")
    a_hh, u_hh = _tail("household_income_amount")
    j = joint(["income", "household_income"])
    j = j[j["income"] != "not_employed"]
    total = j["w"].sum()
    grid_y, grid_f, below = [], [], 0.0
    for b in EARNINGS:
        jb = j[j["income"] == b]
        mass = jb["w"].sum() / total
        if mass == 0:
            continue
        lo_b, top_b, par_b = _bracket(INCOME_BOUNDS[b], a_in, u_in)
        y = np.exp(np.linspace(np.log(lo_b), np.log(top_b), 600))
        f = np.zeros_like(y)
        for r in jb.itertuples(index=False):
            top_h = _bracket(HH_BOUNDS[r.household_income], a_hh, u_hh)[1]
            f += r.w / jb["w"].sum() * _cdf(y, lo_b, min(top_b, top_h), par_b, a_in)
        grid_y.append(y)
        grid_f.append(below + mass * f)
        below += mass
    ys, fs = np.concatenate(grid_y), np.maximum.accumulate(np.concatenate(grid_f))
    bps = [float(np.interp(q, fs, ys)) for q in np.linspace(0, 1, 101)]
    pd.DataFrame({"percentile": range(101), "value": np.round(bps, 1)}).to_csv(paths.CPTS / "income_percentile.csv", index=False)
    log(f"  income_percentile (employed: median {bps[50]:,.0f}, p10 {bps[10]:,.0f}, p90 {bps[90]:,.0f})")


def build_household(log=print) -> None:
    """Household gross monthly income bracket, and home ownership.

    Employed: given living arrangement, own earnings bracket and age. Not employed: given
    labour-force status, age and living arrangement (retirees = outside the labour force at
    65+; compulsory service = outside the labour force at their age). Home ownership: given
    household income, age and living arrangement. Separately for Jews-and-others and for Arabs.
    Someone living with their parents reports their parents' household.
    """
    rows = []
    for pg, t in _per_popgroup_living("household_income_employed").items():
        h = main_effects(t, "hh_income", ["income", "age_s"], ["arrangement"], HH_INCOME, STRENGTH)
        h = h.merge(BANDS, on="age_s").assign(labour_status="employed")
        # a household's income is at least its member's earnings: rule out brackets entirely below
        # them (survey answers occasionally conflict, e.g. household net vs personal gross)
        hh_top = h["hh_income"].map(lambda k: np.inf if HH_BOUNDS[k][1] is None else HH_BOUNDS[k][1])
        own_low = h["income"].map(lambda k: INCOME_BOUNDS[k][0])
        h.loc[(hh_top <= own_low).to_numpy(), "p"] = 0.0
        h["p"] = h["p"] / h.groupby(["arrangement", "income", "age_band"])["p"].transform("sum")
        rows += [h.assign(group=g) for g, p in POPGROUP.items() if p == pg]
    for pg, t in _per_popgroup_living("household_income_by_status").items():
        h = main_effects(t, "hh_income", ["employment", "age_s"], ["arrangement"], HH_INCOME, STRENGTH)
        h = smooth_ages(h, "hh_income").merge(BANDS, on="age_s")
        parts = [
            h[h["employment"] == "unemployed"].assign(labour_status="unemployed"),
            h[(h["employment"] == "not_in_labour_force") & ~h["age_band"].isin(OLD_BANDS)].assign(labour_status="not_in_labour_force"),
            h[(h["employment"] == "not_in_labour_force") & h["age_band"].isin(OLD_BANDS)].assign(labour_status="retired"),
            h[h["employment"] == "not_in_labour_force"].assign(labour_status="compulsory_service"),
        ]
        h = pd.concat(parts, ignore_index=True).assign(income="not_employed")
        rows += [h.assign(group=g) for g, p in POPGROUP.items() if p == pg]
    df = pd.concat(rows, ignore_index=True).rename(columns={"arrangement": "living_arrangement", "hh_income": "value"})
    write_cpt(df[["labour_status", "group", "living_arrangement", "income", "age_band", "value", "p"]], "household_income")
    _write_bounds(HH_BOUNDS, "household_income", "household_income_amount")

    ho = {pg: main_effects(t, "owns_home", ["hh_income", "age_s"], ["arrangement"], ["yes", "no"], STRENGTH)
          for pg, t in _per_popgroup_living("home_ownership").items()}
    df = _to_groups(ho).merge(BANDS, on="age_s").rename(columns={"hh_income": "household_income", "arrangement": "living_arrangement", "owns_home": "value"})
    write_cpt(df[["group", "household_income", "age_band", "living_arrangement", "value", "p"]], "home_ownership")
    log("  household_income, household_income_amount, home_ownership")


def _bracket(bounds: tuple[float, float | None], alpha: float, u_max: float) -> tuple[float, float, bool]:
    """(lo, top, pareto) of a bracket's within-bracket distribution, as the sampler draws it."""
    lo, hi = bounds
    return (lo, lo * (1 - u_max) ** (-1 / alpha), True) if hi is None else (lo, hi, False)


def _cdf(y: np.ndarray, lo: float, top: float, pareto: bool, alpha: float) -> np.ndarray:
    """CDF on [lo, top] of the truncated within-bracket distribution (sample.ts intervalCdf)."""
    y = np.clip(y, lo, top)
    if top <= lo:
        return np.ones_like(y)
    if pareto:
        return (1 - (lo / y) ** alpha) / (1 - (lo / top) ** alpha)
    if lo > 0:
        return (np.log(y) - np.log(lo)) / (np.log(top) - np.log(lo))
    return (y - lo) / (top - lo)


def _quantile(v: np.ndarray, lo: float, top: float, pareto: bool, alpha: float) -> np.ndarray:
    if top <= lo:
        return np.full_like(v, lo)
    if pareto:
        return lo * (1 - v * (1 - (lo / top) ** alpha)) ** (-1 / alpha)
    if lo > 0:
        return np.exp(np.log(lo) + v * (np.log(top) - np.log(lo)))
    return lo + v * (top - lo)


def build_household_percentile(log=print) -> None:
    """Household income at percentiles 0..100 among all adult citizens in the model
    (person-weighted: the rank of the household each adult lives in).

    Within a household bracket the amount is not simply log-uniform: for an employed person it
    is floored at their own earnings (which are capped at the bracket's top). The CDF is
    therefore computed as the exact mixture - people not employed, plus, for each earnings
    bracket, the floored distribution averaged over the person's earnings (64-point quadrature) -
    and inverted numerically.
    """
    a_hh, u_hh = _tail("household_income_amount")
    a_in, u_in = _tail("income_amount")
    j = joint(["household_income", "income"])
    total = j["w"].sum()
    grid_y: list[np.ndarray] = []
    grid_f: list[np.ndarray] = []
    below = 0.0
    for h in HH_INCOME:
        jh = j[j["household_income"] == h]
        mass = jh["w"].sum() / total
        if mass == 0:
            continue
        lo_h, top_h, par_h = _bracket(HH_BOUNDS[h], a_hh, u_hh)
        if top_h <= lo_h:  # a point mass (no income): every percentile it spans maps to it
            grid_y.append(np.array([lo_h, lo_h]))
            grid_f.append(np.array([below, below + mass]))
            below += mass
            continue
        y = np.exp(np.linspace(np.log(lo_h), np.log(top_h), 600))
        f = np.zeros_like(y)
        for r in jh.itertuples(index=False):
            share = r.w / jh["w"].sum()
            if r.income == "not_employed":
                f += share * _cdf(y, lo_h, top_h, par_h, a_hh)
                continue
            lo_b, top_b, par_b = _bracket(INCOME_BOUNDS[r.income], a_in, u_in)
            xs = _quantile((np.arange(64) + 0.5) / 64, lo_b, min(top_b, top_h), par_b, a_in)
            for x in xs:
                floor = min(max(lo_h, x), top_h)
                f += share / 64 * np.where(y >= floor, _cdf(y, floor, top_h, par_h, a_hh), 0.0)
        grid_y.append(y)
        grid_f.append(below + mass * f)
        below += mass
    ys, fs = np.concatenate(grid_y), np.concatenate(grid_f)
    fs = np.maximum.accumulate(fs)
    bps = [float(np.interp(q, fs, ys)) for q in np.linspace(0, 1, 101)]
    pd.DataFrame({"percentile": range(101), "value": np.round(bps, 1)}).to_csv(paths.CPTS / "household_income_percentile.csv", index=False)
    log(f"  household_income_percentile (median {bps[50]:,.0f}, p10 {bps[10]:,.0f}, p90 {bps[90]:,.0f})")
