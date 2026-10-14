"""Discovery helper: list the tables of a CBS Statistical Abstract chapter with their titles.

The CBS publication pages load their table lists with JavaScript, but the Excel files
live at predictable URLs:
    https://www.cbs.gov.il/he/publications/doclib/<year>/<chapter folder>/<prefix>_<NN>.xlsx
This probes NN = 01..max, downloads what exists into data/raw/_discovery/ (gitignored)
and prints each table's title so the ones we need can be pinned in sources.yaml.
"""

from __future__ import annotations

import io

import httpx
import openpyxl

from .. import paths
from .base import HEADERS, throttle

BASE = "https://www.cbs.gov.il/he/publications/doclib"


def _title(content: bytes) -> str:
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    texts: list[str] = []
    for row in ws.iter_rows(min_row=1, max_row=6, values_only=True):
        for v in row:
            if isinstance(v, str) and len(v.strip()) > 8:
                texts.append(" ".join(v.split()))
    return " | ".join(texts[:3])


def discover(year: int, folder: str, prefix: str, max_n: int = 40, log=print) -> list[tuple[str, str]]:
    out_dir = paths.RAW / "_discovery" / str(year) / folder
    out_dir.mkdir(parents=True, exist_ok=True)
    found: list[tuple[str, str]] = []
    with httpx.Client(headers=HEADERS, follow_redirects=True, timeout=60) as client:
        for n in range(1, max_n + 1):
            name = f"{prefix}_{n:02d}.xlsx"
            url = f"{BASE}/{year}/{folder}/{name}"
            throttle(url)
            try:
                resp = client.get(url)
            except httpx.HTTPError:
                continue
            if resp.status_code != 200 or not resp.content.startswith(b"PK"):
                continue
            (out_dir / name).write_bytes(resp.content)
            title = _title(resp.content)
            found.append((url, title))
            log(f"{name}\t{title}")
    return found
